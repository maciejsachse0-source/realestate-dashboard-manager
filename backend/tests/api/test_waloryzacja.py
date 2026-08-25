"""Waloryzacja roczna przez HTTP.

Kryterium akceptacji etapu E8: waloryzacja na zestawie testowym daje kwoty
**co do grosza** zgodne z ręcznym wyliczeniem.
"""

from datetime import date
from decimal import Decimal
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from najem.domena.slowniki import (
    BazaOkresuNajmu,
    OperacjaAudytu,
    RodzajKwoty,
    RodzajWskaznika,
    RodzajZabezpieczenia,
    StatusOkresuNajmu,
    StatusWeryfikacji,
    StatusZabezpieczenia,
    TypLokalu,
    TypWartosci,
    TypZdarzenia,
)
from najem.modele import (
    Budynek,
    LogAudytu,
    Lokal,
    Najemca,
    OkresNajmu,
    ParametrWartosc,
    Zabezpieczenie,
    Zdarzenie,
)

pytestmark = pytest.mark.integracja

ROK = 2027
WSKAZNIK = "3.7"


def umowa_z_czynszem(
    baza: Session,
    budynek: Budynek,
    oznaczenie: str,
    czynsz: str,
    *,
    podlega: bool = True,
    miesiac: int | None = 1,
    rodzaj_wskaznika: RodzajWskaznika | None = RodzajWskaznika.GUS_ROK_DO_ROKU,
    stala_stawka: Decimal | None = None,
    pierwsza_waloryzacja: date | None = None,
    rodzaj_kwoty: RodzajKwoty = RodzajKwoty.NETTO,
    waluta: str = "PLN",
    data_przekazania: date = date(2024, 1, 1),
    okres_miesiace: int = 120,
    data_zakonczenia_faktyczna: date | None = None,
    status: StatusOkresuNajmu = StatusOkresuNajmu.AKTYWNA,
    czynsz_od: date = date(2024, 1, 1),
    status_czynszu: StatusWeryfikacji = StatusWeryfikacji.ZATWIERDZONA,
) -> OkresNajmu:
    lokal = Lokal(budynek_id=budynek.id, oznaczenie=oznaczenie, typ=TypLokalu.BIUROWY)
    baza.add(lokal)
    baza.flush()

    najemca = Najemca(nazwa_pelna=f"Najemca {oznaczenie}")
    baza.add(najemca)
    baza.flush()

    okres = OkresNajmu(
        lokal_id=lokal.id,
        najemca_id=najemca.id,
        data_przekazania=data_przekazania,
        bazuje_na_dacie=BazaOkresuNajmu.DATA_PRZEKAZANIA,
        okres_zawarcia_miesiace=okres_miesiace,
        data_zakonczenia_faktyczna=data_zakonczenia_faktyczna,
        status=status,
        waloryzacja_podlega=podlega,
        waloryzacja_miesiac=miesiac,
        waloryzacja_rodzaj_wskaznika=rodzaj_wskaznika,
        waloryzacja_stala_stawka=stala_stawka,
        waloryzacja_pierwsza_data=pierwsza_waloryzacja,
    )
    baza.add(okres)
    baza.flush()

    baza.add(
        ParametrWartosc(
            okres_najmu_id=okres.id,
            klucz="czynsz_podstawowy",
            typ_wartosci=TypWartosci.KWOTA,
            wartosc_kwota=Decimal(czynsz),
            wartosc_waluta=waluta,
            wartosc_rodzaj_kwoty=rodzaj_kwoty,
            wartosc_stawka_vat=Decimal("23.00") if rodzaj_kwoty is RodzajKwoty.NETTO else None,
            obowiazuje_od=czynsz_od,
            status_weryfikacji=status_czynszu,
        )
    )
    baza.flush()
    return okres


@pytest.fixture
def wskaznik(klient_zarzadca: TestClient) -> None:
    klient_zarzadca.post(
        "/api/v1/waloryzacja/wskazniki",
        json={"rok": ROK, "rodzaj": "gus_rok_do_roku", "wartosc_procent": WSKAZNIK},
    )


class TestWskazniki:
    def test_wprowadzenie_i_odczyt(self, klient_zarzadca: TestClient) -> None:
        odpowiedz = klient_zarzadca.post(
            "/api/v1/waloryzacja/wskazniki",
            json={
                "rok": ROK,
                "rodzaj": "gus_rok_do_roku",
                "wartosc_procent": "3.7",
                "data_publikacji": "2027-01-15",
            },
        )
        assert odpowiedz.status_code == 201
        # Kolumna ma NUMERIC(5,2), więc 3,7 wraca jako 3,70. Obie postacie
        # dają ten sam mnożnik, bo 3.70/100 to nadal 0.037.
        assert odpowiedz.json()["wartosc_procent"] == "3.70"

        lista = klient_zarzadca.get(f"/api/v1/waloryzacja/wskazniki?rok={ROK}").json()
        assert lista["wszystkich"] == 1
        assert len(lista["pozycje"]) == 1

    def test_odpowiedz_zapisu_zgadza_sie_z_odczytem(self, klient_zarzadca: TestClient) -> None:
        """Sesja ma expire_on_commit=False, więc bez odświeżenia POST oddawałby
        to, co przyszło w żądaniu, a GET to, co jest w bazie. Ta sama wartość
        pokazywana na dwa sposoby to problem, nawet jeśli liczbowo jest równa.
        """
        zapis = klient_zarzadca.post(
            "/api/v1/waloryzacja/wskazniki",
            json={"rok": ROK, "rodzaj": "gus_rok_do_roku", "wartosc_procent": "5"},
        ).json()
        odczyt = klient_zarzadca.get(f"/api/v1/waloryzacja/wskazniki?rok={ROK}").json()["pozycje"][
            0
        ]

        assert zapis["wartosc_procent"] == odczyt["wartosc_procent"] == "5.00"

    def test_powtorne_wprowadzenie_jest_odrzucane(
        self, klient_zarzadca: TestClient, wskaznik: None
    ) -> None:
        """GUS nie publikuje wskaźnika dwa razy. Cicha podmiana byłaby gorsza
        niż błąd: przeliczyłaby czynsze inaczej, niż zapowiadał podgląd.
        """
        odpowiedz = klient_zarzadca.post(
            "/api/v1/waloryzacja/wskazniki",
            json={"rok": ROK, "rodzaj": "gus_rok_do_roku", "wartosc_procent": "9.9"},
        )
        assert odpowiedz.status_code == 409
        assert "już wprowadzony" in odpowiedz.json()["detail"]

    def test_podglad_nie_moze_wprowadzac(self, klient_podglad: TestClient) -> None:
        assert (
            klient_podglad.post(
                "/api/v1/waloryzacja/wskazniki",
                json={"rok": ROK, "rodzaj": "gus_rok_do_roku", "wartosc_procent": "3.7"},
            ).status_code
            == 403
        )

    def test_wskaznik_ujemny_jest_dozwolony(self, klient_zarzadca: TestClient) -> None:
        """Deflacja jest rzadka, ale wskaźnik GUS potrafi być ujemny."""
        odpowiedz = klient_zarzadca.post(
            "/api/v1/waloryzacja/wskazniki",
            json={"rok": ROK, "rodzaj": "gus_srednioroczny", "wartosc_procent": "-1.2"},
        )
        assert odpowiedz.status_code == 201


class TestPrzebieg:
    def test_kwoty_co_do_grosza(
        self, klient_zarzadca: TestClient, baza: Session, budynek_api: Budynek, wskaznik: None
    ) -> None:
        """Kryterium akceptacji E8. Ręczne wyliczenie:

        9500,00 razy 1,037 = 9851,50
        4321,99 razy 1,037 = 4481,90363 → 4481,90
        333,33  razy 1,037 = 345,663... → 345,66
        """
        for oznaczenie, czynsz in (
            ("A/01", "9500.00"),
            ("A/02", "4321.99"),
            ("A/03", "333.33"),
        ):
            umowa_z_czynszem(baza, budynek_api, oznaczenie, czynsz)

        wynik = klient_zarzadca.get(f"/api/v1/waloryzacja/przebieg?rok={ROK}").json()
        kwoty = {p["oznaczenie_lokalu"]: p["kwota_nowa"] for p in wynik["objete"]}

        assert kwoty["A/01"] == "9851.50"
        assert kwoty["A/02"] == "4481.90"
        assert kwoty["A/03"] == "345.66"

    def test_suma_zmian(
        self, klient_zarzadca: TestClient, baza: Session, budynek_api: Budynek, wskaznik: None
    ) -> None:
        umowa_z_czynszem(baza, budynek_api, "A/01", "10000.00")
        umowa_z_czynszem(baza, budynek_api, "A/02", "20000.00")

        wynik = klient_zarzadca.get(f"/api/v1/waloryzacja/przebieg?rok={ROK}").json()
        assert wynik["sumy"] == [
            {
                "waluta": "PLN",
                "umow": 2,
                "przed": "30000.00",
                "po": "31110.00",
                "roznica": "1110.00",
            }
        ]

    def test_walut_nie_dodajemy_do_siebie(
        self, klient_zarzadca: TestClient, baza: Session, budynek_api: Budynek, wskaznik: None
    ) -> None:
        """Umowa w EUR i umowa w PLN nie mają wspólnej sumy. Dodanie ich do siebie
        dałoby liczbę, która wygląda na pieniądze i nie znaczy nic.
        """
        umowa_z_czynszem(baza, budynek_api, "A/01", "10000.00")
        umowa_z_czynszem(baza, budynek_api, "A/02", "2000.00", waluta="EUR")

        wynik = klient_zarzadca.get(f"/api/v1/waloryzacja/przebieg?rok={ROK}").json()
        assert [(s["waluta"], s["przed"], s["po"]) for s in wynik["sumy"]] == [
            ("EUR", "2000.00", "2074.00"),
            ("PLN", "10000.00", "10370.00"),
        ]

    def test_suma_tylko_zaznaczonych(
        self, klient_zarzadca: TestClient, baza: Session, budynek_api: Budynek, wskaznik: None
    ) -> None:
        """Podgląd przed zatwierdzeniem: odznaczenie umowy musi zmienić sumę,
        inaczej użytkownik zatwierdza co innego, niż widzi.
        """
        pierwsza = umowa_z_czynszem(baza, budynek_api, "A/01", "10000.00")
        umowa_z_czynszem(baza, budynek_api, "A/02", "20000.00")

        odpowiedz = klient_zarzadca.post(
            "/api/v1/waloryzacja/podsumowanie",
            json={"rok": ROK, "okresy_najmu": [pierwsza.id]},
        )
        assert odpowiedz.status_code == 200
        assert odpowiedz.json() == [
            {
                "waluta": "PLN",
                "umow": 1,
                "przed": "10000.00",
                "po": "10370.00",
                "roznica": "370.00",
            }
        ]

    def test_suma_bez_zaznaczenia_jest_pusta(
        self, klient_zarzadca: TestClient, baza: Session, budynek_api: Budynek, wskaznik: None
    ) -> None:
        umowa_z_czynszem(baza, budynek_api, "A/01", "10000.00")

        odpowiedz = klient_zarzadca.post(
            "/api/v1/waloryzacja/podsumowanie",
            json={"rok": ROK, "okresy_najmu": []},
        )
        assert odpowiedz.status_code == 200
        assert odpowiedz.json() == []

    def test_podglad_niczego_nie_zapisuje(
        self, klient_zarzadca: TestClient, baza: Session, budynek_api: Budynek, wskaznik: None
    ) -> None:
        umowa_z_czynszem(baza, budynek_api, "A/01", "9500.00")
        przed = len(baza.scalars(select(ParametrWartosc)).all())

        klient_zarzadca.get(f"/api/v1/waloryzacja/przebieg?rok={ROK}")
        assert len(baza.scalars(select(ParametrWartosc)).all()) == przed

    def test_waloryzacja_zachowuje_postac_kwoty(
        self, klient_zarzadca: TestClient, baza: Session, budynek_api: Budynek, wskaznik: None
    ) -> None:
        umowa_z_czynszem(baza, budynek_api, "A/01", "1000.00", rodzaj_kwoty=RodzajKwoty.NETTO)
        umowa_z_czynszem(baza, budynek_api, "A/02", "1000.00", rodzaj_kwoty=RodzajKwoty.BRUTTO)

        wynik = klient_zarzadca.get(f"/api/v1/waloryzacja/przebieg?rok={ROK}").json()
        rodzaje = {p["oznaczenie_lokalu"]: p["rodzaj_kwoty"] for p in wynik["objete"]}
        assert rodzaje["A/01"] == "netto"
        assert rodzaje["A/02"] == "brutto"


class TestWylaczenia:
    def test_umowa_bez_waloryzacji_z_powodem(
        self, klient_zarzadca: TestClient, baza: Session, budynek_api: Budynek, wskaznik: None
    ) -> None:
        """Ekran pokazuje umowy wyłączone razem z powodem wyłączenia."""
        umowa_z_czynszem(baza, budynek_api, "A/01", "1000.00", podlega=False)

        wynik = klient_zarzadca.get(f"/api/v1/waloryzacja/przebieg?rok={ROK}").json()
        assert wynik["objete"] == []
        assert len(wynik["wylaczone"]) == 1
        assert "nie podlega" in wynik["wylaczone"][0]["powod_wylaczenia"]

    def test_brak_wskaznika_dla_rodzaju(
        self, klient_zarzadca: TestClient, baza: Session, budynek_api: Budynek, wskaznik: None
    ) -> None:
        umowa_z_czynszem(
            baza,
            budynek_api,
            "A/01",
            "1000.00",
            rodzaj_wskaznika=RodzajWskaznika.GUS_SREDNIOROCZNY,
        )
        wynik = klient_zarzadca.get(f"/api/v1/waloryzacja/przebieg?rok={ROK}").json()
        assert "Brak wprowadzonego wskaźnika" in wynik["wylaczone"][0]["powod_wylaczenia"]

    def test_umowa_bez_zatwierdzonego_czynszu(
        self, klient_zarzadca: TestClient, baza: Session, budynek_api: Budynek, wskaznik: None
    ) -> None:
        """Decyzja D5: brak czynszu nie waloryzuje się do zera."""
        okres = umowa_z_czynszem(baza, budynek_api, "A/01", "1000.00")
        parametr = baza.scalars(
            select(ParametrWartosc).where(ParametrWartosc.okres_najmu_id == okres.id)
        ).one()
        parametr.status_weryfikacji = StatusWeryfikacji.ZAPROPONOWANA
        baza.flush()

        wynik = klient_zarzadca.get(f"/api/v1/waloryzacja/przebieg?rok={ROK}").json()
        assert wynik["objete"] == []
        assert "nieustalony" in wynik["wylaczone"][0]["powod_wylaczenia"]

    def test_stala_stawka_z_umowy_dziala_bez_wskaznika_gus(
        self, klient_zarzadca: TestClient, baza: Session, budynek_api: Budynek
    ) -> None:
        umowa_z_czynszem(
            baza,
            budynek_api,
            "A/01",
            "1000.00",
            rodzaj_wskaznika=RodzajWskaznika.STALA_STAWKA,
            stala_stawka=Decimal("5"),
        )
        wynik = klient_zarzadca.get(f"/api/v1/waloryzacja/przebieg?rok={ROK}").json()
        assert wynik["objete"][0]["kwota_nowa"] == "1050.00"

    def test_umowa_przed_pierwsza_waloryzacja(
        self, klient_zarzadca: TestClient, baza: Session, budynek_api: Budynek, wskaznik: None
    ) -> None:
        umowa_z_czynszem(
            baza,
            budynek_api,
            "A/01",
            "1000.00",
            pierwsza_waloryzacja=date(2029, 1, 1),
        )
        wynik = klient_zarzadca.get(f"/api/v1/waloryzacja/przebieg?rok={ROK}").json()
        assert "2029" in wynik["wylaczone"][0]["powod_wylaczenia"]


class TestZatwierdzenie:
    def test_zapisuje_nowy_czynsz_od_miesiaca_waloryzacji(
        self, klient_zarzadca: TestClient, baza: Session, budynek_api: Budynek, wskaznik: None
    ) -> None:
        okres = umowa_z_czynszem(baza, budynek_api, "A/01", "9500.00")

        odpowiedz = klient_zarzadca.post(
            "/api/v1/waloryzacja/zatwierdz",
            json={"rok": ROK, "okresy_najmu": [okres.id]},
        )
        assert odpowiedz.status_code == 200
        assert odpowiedz.json()["umow_zwaloryzowanych"] == 1

        stan = klient_zarzadca.get(
            f"/api/v1/lokale/{okres.lokal_id}/stan?na_dzien={ROK}-01-01"
        ).json()
        assert stan["parametry"]["czynsz_podstawowy"]["wartosc"] == "9851.50"

        # Dzień wcześniej obowiązuje jeszcze stara kwota (decyzja D2).
        przed = klient_zarzadca.get(
            f"/api/v1/lokale/{okres.lokal_id}/stan?na_dzien={ROK - 1}-12-31"
        ).json()
        assert przed["parametry"]["czynsz_podstawowy"]["wartosc"] == "9500.00"

    def test_nowa_wartosc_jest_zatwierdzona(
        self, klient_zarzadca: TestClient, baza: Session, budynek_api: Budynek, wskaznik: None
    ) -> None:
        """Człowiek zatwierdził ją na ekranie, oglądając kwotę przed i po."""
        okres = umowa_z_czynszem(baza, budynek_api, "A/01", "9500.00")
        klient_zarzadca.post(
            "/api/v1/waloryzacja/zatwierdz", json={"rok": ROK, "okresy_najmu": [okres.id]}
        )

        nowy = baza.scalars(
            select(ParametrWartosc).where(
                ParametrWartosc.okres_najmu_id == okres.id,
                ParametrWartosc.obowiazuje_od == date(ROK, 1, 1),
            )
        ).one()
        assert nowy.status_weryfikacji is StatusWeryfikacji.ZATWIERDZONA
        assert nowy.zatwierdzil_uzytkownik_id is not None
        assert nowy.uwagi is not None
        assert "Waloryzacja" in nowy.uwagi

    def test_zatwierdzamy_tylko_wybrane_umowy(
        self, klient_zarzadca: TestClient, baza: Session, budynek_api: Budynek, wskaznik: None
    ) -> None:
        pierwsza = umowa_z_czynszem(baza, budynek_api, "A/01", "1000.00")
        druga = umowa_z_czynszem(baza, budynek_api, "A/02", "2000.00")

        klient_zarzadca.post(
            "/api/v1/waloryzacja/zatwierdz", json={"rok": ROK, "okresy_najmu": [pierwsza.id]}
        )

        nowe = baza.scalars(
            select(ParametrWartosc).where(ParametrWartosc.obowiazuje_od == date(ROK, 1, 1))
        ).all()
        assert len(nowe) == 1
        assert nowe[0].okres_najmu_id == pierwsza.id
        assert druga.id != nowe[0].okres_najmu_id

    def test_drugi_przebieg_nie_waloryzuje_ponownie(
        self, klient_zarzadca: TestClient, baza: Session, budynek_api: Budynek, wskaznik: None
    ) -> None:
        """Bez tego zabezpieczenia drugie kliknięcie podniosłoby czynsz
        o kolejne 3,7% od już podniesionej kwoty.
        """
        okres = umowa_z_czynszem(baza, budynek_api, "A/01", "9500.00")
        klient_zarzadca.post(
            "/api/v1/waloryzacja/zatwierdz", json={"rok": ROK, "okresy_najmu": [okres.id]}
        )

        wynik = klient_zarzadca.get(f"/api/v1/waloryzacja/przebieg?rok={ROK}").json()
        assert wynik["objete"] == []
        assert "waloryzowana w przebiegu" in wynik["wylaczone"][0]["powod_wylaczenia"]

        powtorka = klient_zarzadca.post(
            "/api/v1/waloryzacja/zatwierdz", json={"rok": ROK, "okresy_najmu": [okres.id]}
        )
        assert powtorka.status_code == 409

    def test_weksel_do_przeliczenia(
        self, klient_zarzadca: TestClient, baza: Session, budynek_api: Budynek, wskaznik: None
    ) -> None:
        """Reguła R5: zabezpieczenie wyliczane z czynszu trzeba przeliczyć."""
        okres = umowa_z_czynszem(baza, budynek_api, "A/01", "9500.00")
        baza.add(
            Zabezpieczenie(
                okres_najmu_id=okres.id,
                rodzaj=RodzajZabezpieczenia.WEKSEL,
                status=StatusZabezpieczenia.DOSTARCZONE,
                sposob_wyliczenia="czterokrotność czynszu",
            )
        )
        baza.flush()

        odpowiedz = klient_zarzadca.post(
            "/api/v1/waloryzacja/zatwierdz", json={"rok": ROK, "okresy_najmu": [okres.id]}
        )
        assert odpowiedz.json()["zdarzen_o_wekslach"] == 1

        zdarzenie = baza.scalars(
            select(Zdarzenie).where(Zdarzenie.typ == TypZdarzenia.WEKSEL_DO_PRZELICZENIA)
        ).one()
        assert "czterokrotność czynszu" in zdarzenie.tresc
        assert zdarzenie.data_zdarzenia == date(ROK, 1, 1)

    def test_zabezpieczenie_o_stalej_wartosci_nie_wymaga_przeliczenia(
        self, klient_zarzadca: TestClient, baza: Session, budynek_api: Budynek, wskaznik: None
    ) -> None:
        okres = umowa_z_czynszem(baza, budynek_api, "A/01", "9500.00")
        baza.add(
            Zabezpieczenie(
                okres_najmu_id=okres.id,
                rodzaj=RodzajZabezpieczenia.KAUCJA,
                status=StatusZabezpieczenia.DOSTARCZONE,
                wymagana_wartosc=Decimal("20000.00"),
                wymagana_waluta="PLN",
                wymagana_rodzaj_kwoty=RodzajKwoty.BRUTTO,
            )
        )
        baza.flush()

        odpowiedz = klient_zarzadca.post(
            "/api/v1/waloryzacja/zatwierdz", json={"rok": ROK, "okresy_najmu": [okres.id]}
        )
        assert odpowiedz.json()["zdarzen_o_wekslach"] == 0

    def test_umowa_spoza_przebiegu_jest_odrzucana(
        self, klient_zarzadca: TestClient, baza: Session, budynek_api: Budynek, wskaznik: None
    ) -> None:
        """Lista mogła się zdezaktualizować, gdy ktoś zmienił umowę w międzyczasie."""
        wylaczona = umowa_z_czynszem(baza, budynek_api, "A/01", "1000.00", podlega=False)

        odpowiedz = klient_zarzadca.post(
            "/api/v1/waloryzacja/zatwierdz",
            json={"rok": ROK, "okresy_najmu": [wylaczona.id]},
        )
        assert odpowiedz.status_code == 409
        assert "Odśwież listę" in odpowiedz.json()["detail"]

    def test_podglad_nie_moze_zatwierdzac(
        self, klient_podglad: TestClient, baza: Session, budynek_api: Budynek
    ) -> None:
        okres = umowa_z_czynszem(baza, budynek_api, "A/01", "1000.00")
        assert (
            klient_podglad.post(
                "/api/v1/waloryzacja/zatwierdz", json={"rok": ROK, "okresy_najmu": [okres.id]}
            ).status_code
            == 403
        )


class TestEksport:
    def test_arkusz_z_lista_zmian(
        self, klient_zarzadca: TestClient, baza: Session, budynek_api: Budynek, wskaznik: None
    ) -> None:
        okres = umowa_z_czynszem(baza, budynek_api, "A/01", "9500.00")
        umowa_z_czynszem(baza, budynek_api, "A/02", "1000.00", podlega=False)
        klient_zarzadca.post(
            "/api/v1/waloryzacja/zatwierdz", json={"rok": ROK, "okresy_najmu": [okres.id]}
        )

        odpowiedz = klient_zarzadca.get(f"/api/v1/waloryzacja/eksport?rok={ROK}")
        assert odpowiedz.status_code == 200
        assert "waloryzacja-2027.xlsx" in odpowiedz.headers["Content-Disposition"]

        skoroszyt = load_workbook(BytesIO(odpowiedz.content))
        arkusz = skoroszyt[f"Waloryzacja {ROK}"]
        wiersze = list(arkusz.iter_rows(values_only=True))
        assert wiersze[0][0] == "Budynek i lokal"
        assert wiersze[1][0] == "A/01"
        assert Decimal(str(wiersze[1][2])) == Decimal("9500.00")
        assert Decimal(str(wiersze[1][3])) == Decimal("9851.50")

        # Wyłączone lądują w osobnym arkuszu razem z powodem. Jest tam też
        # A/01, bo po zatwierdzeniu wypada z przebiegu jako już zwaloryzowana.
        powody = {w[0]: w[2] for w in skoroszyt["Wyłączone"].iter_rows(min_row=2, values_only=True)}
        assert "nie podlega" in str(powody["A/02"])
        assert "waloryzowana w przebiegu" in str(powody["A/01"])

    def test_grosze_przezywaja_eksport(
        self, klient_zarzadca: TestClient, baza: Session, budynek_api: Budynek, wskaznik: None
    ) -> None:
        """Kwoty idą do arkusza jako `Decimal`, bez konwersji na `float`
        po drodze — zasada twarda projektu nie robi wyjątku dla eksportów.

        Sam plik XLSX nie ma typu dziesiętnego (liczby to `double`), więc
        odczyt zawsze zwróci `float`. Sprawdzamy więc to, co da się sprawdzić:
        że po drodze nie zgubił się ani jeden grosz.
        """
        okres = umowa_z_czynszem(baza, budynek_api, "A/01", "4321.99")
        klient_zarzadca.post(
            "/api/v1/waloryzacja/zatwierdz", json={"rok": ROK, "okresy_najmu": [okres.id]}
        )

        odpowiedz = klient_zarzadca.get(f"/api/v1/waloryzacja/eksport?rok={ROK}")
        arkusz = load_workbook(BytesIO(odpowiedz.content))[f"Waloryzacja {ROK}"]

        assert Decimal(str(arkusz["C2"].value)) == Decimal("4321.99")
        assert Decimal(str(arkusz["D2"].value)) == Decimal("4481.90")
        assert Decimal(str(arkusz["E2"].value)) == Decimal("159.91")

    def test_arkusz_dziala_po_zatwierdzeniu(
        self, klient_zarzadca: TestClient, baza: Session, budynek_api: Budynek, wskaznik: None
    ) -> None:
        """Pisma do najemców pisze się PO zatwierdzeniu. Gdyby arkusz pokazywał
        tylko niezatwierdzone propozycje, byłby pusty dokładnie wtedy, kiedy
        jest potrzebny.
        """
        okres = umowa_z_czynszem(baza, budynek_api, "A/01", "9500.00")
        klient_zarzadca.post(
            "/api/v1/waloryzacja/zatwierdz", json={"rok": ROK, "okresy_najmu": [okres.id]}
        )

        odpowiedz = klient_zarzadca.get(f"/api/v1/waloryzacja/eksport?rok={ROK}")
        arkusz = load_workbook(BytesIO(odpowiedz.content))[f"Waloryzacja {ROK}"]
        wiersze = list(arkusz.iter_rows(values_only=True))

        assert wiersze[1][0] == "A/01"
        assert Decimal(str(wiersze[1][2])) == Decimal("9500.00")
        assert Decimal(str(wiersze[1][3])) == Decimal("9851.50")
        # Wskaźnik pochodzi ze znacznika zapisanego przy zatwierdzaniu,
        # a nie z odtwarzania go z pary kwot.
        assert Decimal(str(wiersze[1][7])) == Decimal("3.70")

    def test_propozycje_sa_w_osobnym_arkuszu(
        self, klient_zarzadca: TestClient, baza: Session, budynek_api: Budynek, wskaznik: None
    ) -> None:
        """Decyzja D4: wartość niezatwierdzona nie wchodzi do raportu. Z tego
        arkusza ktoś robi korespondencję seryjną do najemców.
        """
        zatwierdzona = umowa_z_czynszem(baza, budynek_api, "A/01", "9500.00")
        umowa_z_czynszem(baza, budynek_api, "A/02", "1000.00")
        klient_zarzadca.post(
            "/api/v1/waloryzacja/zatwierdz",
            json={"rok": ROK, "okresy_najmu": [zatwierdzona.id]},
        )

        odpowiedz = klient_zarzadca.get(f"/api/v1/waloryzacja/eksport?rok={ROK}")
        skoroszyt = load_workbook(BytesIO(odpowiedz.content))

        zatwierdzone = [
            w[0] for w in skoroszyt[f"Waloryzacja {ROK}"].iter_rows(min_row=2, values_only=True)
        ]
        propozycje = [
            w[0]
            for w in skoroszyt["Propozycje niezatwierdzone"].iter_rows(min_row=2, values_only=True)
        ]

        assert zatwierdzone == ["A/01"]
        assert propozycje == ["A/02"]

    def test_bez_propozycji_nie_ma_takiego_arkusza(
        self, klient_zarzadca: TestClient, baza: Session, budynek_api: Budynek, wskaznik: None
    ) -> None:
        okres = umowa_z_czynszem(baza, budynek_api, "A/01", "9500.00")
        klient_zarzadca.post(
            "/api/v1/waloryzacja/zatwierdz", json={"rok": ROK, "okresy_najmu": [okres.id]}
        )

        odpowiedz = klient_zarzadca.get(f"/api/v1/waloryzacja/eksport?rok={ROK}")
        assert (
            "Propozycje niezatwierdzone" not in load_workbook(BytesIO(odpowiedz.content)).sheetnames
        )

    def test_kwoty_maja_format_ksiegowy(
        self, klient_zarzadca: TestClient, baza: Session, budynek_api: Budynek, wskaznik: None
    ) -> None:
        """Arkusz idzie do pism dla najemców. Bez formatu Excel pokazałby
        9851,5 zamiast 9 851,50, a to na piśmie wygląda na pomyłkę.
        """
        okres = umowa_z_czynszem(baza, budynek_api, "A/01", "9500.00")
        klient_zarzadca.post(
            "/api/v1/waloryzacja/zatwierdz", json={"rok": ROK, "okresy_najmu": [okres.id]}
        )

        odpowiedz = klient_zarzadca.get(f"/api/v1/waloryzacja/eksport?rok={ROK}")
        arkusz = load_workbook(BytesIO(odpowiedz.content))[f"Waloryzacja {ROK}"]

        assert arkusz["C2"].number_format == "# ##0.00"  # czynsz przed
        assert arkusz["D2"].number_format == "# ##0.00"  # czynsz po
        # Procent nie jest kwotą i nie dostaje separatora tysięcy.
        assert arkusz["H2"].number_format == "0.00"
        assert arkusz["I2"].number_format == "DD.MM.YYYY"  # obowiązuje od
        assert arkusz["A1"].font.bold
        assert arkusz.freeze_panes == "A2"


class TestPrzypadkiBrzegowe:
    """Przypadki, ktore CLAUDE.md wymienia jako obowiazkowe, plus te,
    ktore wyszly w przegladzie kodu E8.
    """

    def test_luty_w_roku_przestepnym(
        self, klient_zarzadca: TestClient, baza: Session, budynek_api: Budynek
    ) -> None:
        """Waloryzacja od 1 marca bierze czynsz z dnia poprzedniego, czyli
        z 29 lutego. W roku nieprzestepnym tego dnia nie ma.
        """
        klient_zarzadca.post(
            "/api/v1/waloryzacja/wskazniki",
            json={"rok": 2028, "rodzaj": "gus_rok_do_roku", "wartosc_procent": "10"},
        )
        umowa_z_czynszem(baza, budynek_api, "A/01", "1000.00", miesiac=3)

        wynik = klient_zarzadca.get("/api/v1/waloryzacja/przebieg?rok=2028").json()
        pozycja = wynik["objete"][0]

        assert pozycja["obowiazuje_od"] == "2028-03-01"
        assert pozycja["kwota_nowa"] == "1100.00"

    def test_deflacja_daje_ujemna_roznice(
        self, klient_zarzadca: TestClient, baza: Session, budynek_api: Budynek
    ) -> None:
        """Wskaznik ujemny jest dopuszczalny, wiec roznica i suma musza wyjsc
        ujemne. Bez tego testu interfejs pokazywal obnizke jako wzrost.
        """
        klient_zarzadca.post(
            "/api/v1/waloryzacja/wskazniki",
            json={"rok": ROK, "rodzaj": "gus_rok_do_roku", "wartosc_procent": "-2.5"},
        )
        umowa_z_czynszem(baza, budynek_api, "A/01", "10000.00")

        wynik = klient_zarzadca.get(f"/api/v1/waloryzacja/przebieg?rok={ROK}").json()

        assert wynik["objete"][0]["kwota_nowa"] == "9750.00"
        assert wynik["objete"][0]["roznica"] == "-250.00"
        assert wynik["sumy"][0]["roznica"] == "-250.00"

    def test_umowa_konczaca_sie_przed_waloryzacja_wypada(
        self, klient_zarzadca: TestClient, baza: Session, budynek_api: Budynek, wskaznik: None
    ) -> None:
        """Status okresu najmu zmienia czlowiek, wiec umowa zakonczona w czerwcu
        potrafi wisiec w bazie jako aktywna. Podwyzka dla najmu, ktory juz nie
        trwa, trafilaby do pisma dla bylego najemcy.
        """
        umowa_z_czynszem(
            baza,
            budynek_api,
            "A/01",
            "9500.00",
            data_zakonczenia_faktyczna=date(ROK - 1, 12, 31),
            status=StatusOkresuNajmu.WYPOWIEDZIANA,
        )

        wynik = klient_zarzadca.get(f"/api/v1/waloryzacja/przebieg?rok={ROK}").json()

        assert wynik["objete"] == []
        assert "kończy się 31.12.2026" in wynik["wylaczone"][0]["powod_wylaczenia"]

    def test_umowa_zaczynajaca_sie_po_waloryzacji_wypada(
        self, klient_zarzadca: TestClient, baza: Session, budynek_api: Budynek, wskaznik: None
    ) -> None:
        umowa_z_czynszem(
            baza,
            budynek_api,
            "A/01",
            "9500.00",
            data_przekazania=date(ROK, 6, 1),
            czynsz_od=date(ROK, 6, 1),
        )

        wynik = klient_zarzadca.get(f"/api/v1/waloryzacja/przebieg?rok={ROK}").json()

        assert wynik["objete"] == []
        assert "zaczyna się 01.06.2027" in wynik["wylaczone"][0]["powod_wylaczenia"]

    def test_niezatwierdzona_propozycja_nie_blokuje_przebiegu(
        self, klient_zarzadca: TestClient, baza: Session, budynek_api: Budynek, wskaznik: None
    ) -> None:
        """Wczesniej wystarczyla dowolna wersja czynszu wchodzaca w dniu
        waloryzacji, zeby umowa wypadla z komunikatem, ze podwyzka juz byla.
        Wersja niezatwierdzona nie jest podwyzka (decyzja D4).
        """
        okres = umowa_z_czynszem(baza, budynek_api, "A/01", "9500.00")
        baza.add(
            ParametrWartosc(
                okres_najmu_id=okres.id,
                klucz="czynsz_podstawowy",
                typ_wartosci=TypWartosci.KWOTA,
                wartosc_kwota=Decimal("9999.00"),
                wartosc_waluta="PLN",
                wartosc_rodzaj_kwoty=RodzajKwoty.NETTO,
                wartosc_stawka_vat=Decimal("23.00"),
                obowiazuje_od=date(ROK, 1, 1),
                status_weryfikacji=StatusWeryfikacji.ZAPROPONOWANA,
            )
        )
        baza.flush()

        wynik = klient_zarzadca.get(f"/api/v1/waloryzacja/przebieg?rok={ROK}").json()

        assert [p["oznaczenie_lokalu"] for p in wynik["objete"]] == ["A/01"]
        assert wynik["objete"][0]["kwota_nowa"] == "9851.50"

    def test_aneks_w_dniu_waloryzacji_nie_jest_waloryzacja(
        self, klient_zarzadca: TestClient, baza: Session, budynek_api: Budynek, wskaznik: None
    ) -> None:
        """Aneks wchodzacy 1 stycznia wygladal w eksporcie jak waloryzacja,
        z dorobionym wskaznikiem, ktorego nigdy nie bylo.
        """
        okres = umowa_z_czynszem(baza, budynek_api, "A/01", "9500.00")
        baza.add(
            ParametrWartosc(
                okres_najmu_id=okres.id,
                klucz="czynsz_podstawowy",
                typ_wartosci=TypWartosci.KWOTA,
                wartosc_kwota=Decimal("12000.00"),
                wartosc_waluta="PLN",
                wartosc_rodzaj_kwoty=RodzajKwoty.NETTO,
                wartosc_stawka_vat=Decimal("23.00"),
                obowiazuje_od=date(ROK, 1, 1),
                status_weryfikacji=StatusWeryfikacji.ZATWIERDZONA,
                uwagi="Aneks nr 2, nowa stawka wynegocjowana",
            )
        )
        baza.flush()

        odpowiedz = klient_zarzadca.get(f"/api/v1/waloryzacja/eksport?rok={ROK}")
        arkusz = load_workbook(BytesIO(odpowiedz.content))[f"Waloryzacja {ROK}"]

        # Arkusz zatwierdzonych zmian ma sam naglowek: aneks to nie waloryzacja.
        assert arkusz.max_row == 1

    def test_zatwierdzenie_zostawia_slad_w_audycie(
        self, klient_zarzadca: TestClient, baza: Session, budynek_api: Budynek, wskaznik: None
    ) -> None:
        """Regula twarda: kazda zmiana danych zapisuje wpis w log_audytu."""
        okres = umowa_z_czynszem(baza, budynek_api, "A/01", "9500.00")
        klient_zarzadca.post(
            "/api/v1/waloryzacja/zatwierdz", json={"rok": ROK, "okresy_najmu": [okres.id]}
        )

        nowy = baza.scalars(
            select(ParametrWartosc).where(
                ParametrWartosc.okres_najmu_id == okres.id,
                ParametrWartosc.waloryzacja_rok == ROK,
            )
        ).one()
        wpisy = baza.scalars(
            select(LogAudytu).where(
                LogAudytu.tabela == "parametr_wartosc",
                LogAudytu.rekord_id == nowy.id,
            )
        ).all()

        assert len(wpisy) == 1
        assert wpisy[0].operacja is OperacjaAudytu.UTWORZENIE
        assert wpisy[0].uzytkownik_id is not None
        assert wpisy[0].adres_ip is not None


class TestZeroProcent:
    def test_wskaznik_zero_nie_wystawia_zdarzen_o_zabezpieczeniach(
        self, klient_zarzadca: TestClient, baza: Session, budynek_api: Budynek
    ) -> None:
        """Przy 0% czynsz sie nie zmienia, wiec zabezpieczenie nadal pokrywa
        te sama ekspozycje. Alarm bez pokrycia w danych uczy ludzi ignorowania
        alarmow.
        """
        klient_zarzadca.post(
            "/api/v1/waloryzacja/wskazniki",
            json={"rok": ROK, "rodzaj": "gus_rok_do_roku", "wartosc_procent": "0"},
        )
        okres = umowa_z_czynszem(baza, budynek_api, "A/01", "9500.00")
        baza.add(
            Zabezpieczenie(
                okres_najmu_id=okres.id,
                rodzaj=RodzajZabezpieczenia.WEKSEL,
                status=StatusZabezpieczenia.DOSTARCZONE,
                sposob_wyliczenia="czterokrotność czynszu podstawowego",
            )
        )
        baza.flush()

        odpowiedz = klient_zarzadca.post(
            "/api/v1/waloryzacja/zatwierdz", json={"rok": ROK, "okresy_najmu": [okres.id]}
        )

        assert odpowiedz.status_code == 200
        assert odpowiedz.json()["zdarzen_o_wekslach"] == 0
        assert (
            baza.scalars(
                select(Zdarzenie).where(Zdarzenie.typ == TypZdarzenia.WEKSEL_DO_PRZELICZENIA)
            ).all()
            == []
        )
